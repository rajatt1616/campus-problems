"""Build a simple 8-slide CampusComplaint presentation."""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE


INDIGO = RGBColor(0x43, 0x38, 0xCA)
INDIGO_DARK = RGBColor(0x37, 0x30, 0xA3)
LIGHT = RGBColor(0xEE, 0xF2, 0xFF)
INK = RGBColor(0x18, 0x18, 0x1B)
MUTED = RGBColor(0x71, 0x71, 0x7A)
BORDER = RGBColor(0xE4, 0xE4, 0xE7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAGE = RGBColor(0xFA, 0xFA, 0xFA)

FONT = "Segoe UI"
SLIDE_W = 13.333
SLIDE_H = 7.5
FOOTER = "CampusComplaint  ·  Campus Grievance Management System"


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background
    background.fill.solid()
    background.fill.fore_color.rgb = WHITE
    return slide


def add_text(slide, text, x, y, w, h, size=16, bold=False, color=INK,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def add_bullets(slide, items, x, y, w, h, size=17, space_after=12):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0

    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(space_after)
        p.line_spacing = 1.12

        lead, detail = (item if isinstance(item, tuple) else (None, item))

        bullet = p.add_run()
        bullet.text = "\u2022  "
        bullet.font.name = FONT
        bullet.font.size = Pt(size)
        bullet.font.bold = True
        bullet.font.color.rgb = INDIGO

        if lead:
            run = p.add_run()
            run.text = lead
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.bold = True
            run.font.color.rgb = INK

            run = p.add_run()
            run.text = " " + detail
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.color.rgb = INK
        else:
            run = p.add_run()
            run.text = detail
            run.font.name = FONT
            run.font.size = Pt(size)
            run.font.color.rgb = INK
    return box


def add_card(slide, x, y, w, h, fill=WHITE, line=BORDER, radius=0.04):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.adjustments[0] = radius
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(1)
    shape.shadow.inherit = False
    return shape


def add_title(slide, title, subtitle=None, number=None):
    add_text(slide, title, 0.7, 0.45, 11.9, 0.55, size=30, bold=True, color=INK)
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.1), Inches(1.0), Inches(0.06)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = INDIGO
    bar.line.fill.background()
    bar.shadow.inherit = False
    if subtitle:
        add_text(slide, subtitle, 0.7, 1.28, 11.9, 0.35, size=13, color=MUTED)


def add_footer(slide, number):
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(7.02), Inches(11.93), Inches(0.01)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = BORDER
    line.line.fill.background()
    line.shadow.inherit = False
    add_text(slide, FOOTER, 0.7, 7.12, 9.0, 0.25, size=9, color=MUTED)
    add_text(slide, str(number), 11.6, 7.12, 1.03, 0.25, size=9,
             color=MUTED, align=PP_ALIGN.RIGHT)


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def build():
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    # ---------------------------------------------------------------- slide 1
    slide = new_slide(prs)

    logo = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.17), Inches(1.15), Inches(1.0), Inches(1.0)
    )
    logo.adjustments[0] = 0.18
    logo.fill.solid()
    logo.fill.fore_color.rgb = INDIGO
    logo.line.fill.background()
    logo.shadow.inherit = False
    tf = logo.text_frame
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "CC"
    r.font.name = FONT
    r.font.size = Pt(26)
    r.font.bold = True
    r.font.color.rgb = WHITE

    add_text(slide, "CampusComplaint", 0.7, 2.45, 11.93, 0.9,
             size=46, bold=True, color=INK, align=PP_ALIGN.CENTER)
    add_text(slide, "File.  Assign.  Resolve.  Verify.", 0.7, 3.4, 11.93, 0.45,
             size=20, bold=True, color=INDIGO, align=PP_ALIGN.CENTER)

    divider = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(6.17), Inches(4.05), Inches(1.0), Inches(0.05)
    )
    divider.fill.solid()
    divider.fill.fore_color.rgb = BORDER
    divider.line.fill.background()
    divider.shadow.inherit = False

    add_text(slide, "Campus Grievance Management System", 0.7, 4.3, 11.93, 0.4,
             size=16, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "Django Web Application  ·  Project Presentation",
             0.7, 4.75, 11.93, 0.35, size=13, color=MUTED, align=PP_ALIGN.CENTER)

    add_notes(
        slide,
        "CampusComplaint is a Django web app that manages campus complaints from "
        "filing through verified closure. The tagline sums up the four stages: "
        "file, assign, resolve, verify.",
    )

    # ---------------------------------------------------------------- slide 2
    slide = new_slide(prs)
    add_title(slide, "What It Is", "One tracked workflow for campus complaints")

    left = add_card(slide, 0.7, 1.75, 5.85, 4.25, fill=WHITE)
    add_text(slide, "The Problem", 1.0, 1.98, 5.2, 0.4, size=18, bold=True, color=INK)
    add_bullets(
        slide,
        [
            "Complaints live in paper forms, chats, and verbal requests",
            "Issues get lost with no clear owner",
            "No record of what changed, or when",
            "Students cannot confirm a fix",
        ],
        1.0, 2.5, 5.25, 3.2, size=15, space_after=14,
    )

    right = add_card(slide, 6.78, 1.75, 5.85, 4.25, fill=LIGHT, line=LIGHT)
    add_text(slide, "The Solution", 7.08, 1.98, 5.2, 0.4, size=18, bold=True, color=INDIGO)
    add_bullets(
        slide,
        [
            "One web app for the entire complaint journey",
            "Every complaint has an ID, an owner, and a status",
            "Full status history and remarks on every issue",
            "Students verify and close the resolution",
        ],
        7.08, 2.5, 5.25, 3.2, size=15, space_after=14,
    )

    add_text(
        slide,
        "CampusComplaint turns scattered complaints into a tracked, accountable workflow.",
        0.7, 6.25, 11.93, 0.4, size=15, bold=True, color=INDIGO,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, 2)
    add_notes(
        slide,
        "The problem is not that complaints are rare; it is that they are scattered. "
        "CampusComplaint gives each one a single tracked path with a clear owner.",
    )

    # ---------------------------------------------------------------- slide 3
    slide = new_slide(prs)
    add_title(slide, "Features", "Everything needed to run a complaint end to end")

    add_bullets(
        slide,
        [
            ("Filing \u2014", "students submit a complaint with category and priority"),
            ("Triage \u2014", "admins see all complaints and assign them to a department"),
            ("Resolution \u2014", "departments start work and resolve with a note"),
        ],
        0.7, 1.85, 5.85, 4.6, size=16, space_after=24,
    )
    add_bullets(
        slide,
        [
            ("Verification \u2014", "students accept, reject with a reason, or reopen"),
            ("Visibility \u2014", "remarks, status timeline, notifications, and search"),
            ("Admin tools \u2014", "JSON export and a Django admin panel"),
        ],
        6.78, 1.85, 5.85, 4.6, size=16, space_after=24,
    )

    add_card(slide, 0.7, 5.35, 11.93, 0.95, fill=LIGHT, line=LIGHT)
    add_text(
        slide,
        "Every complaint gets an automatic ID such as CMP-0001, so it can be tracked and referenced.",
        0.95, 5.55, 11.4, 0.5, size=14, color=INK,
    )
    add_footer(slide, 3)
    add_notes(
        slide,
        "Features follow the life of a complaint: filing, triage, resolution, "
        "verification, and visibility. Auto IDs make each issue easy to reference.",
    )

    # ---------------------------------------------------------------- slide 4
    slide = new_slide(prs)
    add_title(slide, "Roles and Permissions", "Three roles, each with a clear scope")

    rows, cols = 4, 2
    table_shape = slide.shapes.add_table(
        rows, cols, Inches(0.7), Inches(1.75), Inches(11.93), Inches(4.35)
    )
    table = table_shape.table
    table.columns[0].width = Inches(2.45)
    table.columns[1].width = Inches(9.48)
    table.rows[0].height = Inches(0.6)
    for r in range(1, rows):
        table.rows[r].height = Inches(1.25)

    data = [
        ("Role", "What they can do"),
        ("Student", "File complaints  ·  edit or delete while filed  ·  accept, reject, or reopen  ·  view own complaints"),
        ("Admin", "View all complaints  ·  assign to departments  ·  delete  ·  export JSON  ·  use Django admin"),
        ("Department", "View assigned complaints  ·  start work  ·  mark resolved with a note  ·  add remarks"),
    ]

    for r in range(rows):
        for c in range(cols):
            cell = table.cell(r, c)
            cell.margin_left = Inches(0.18)
            cell.margin_right = Inches(0.18)
            cell.margin_top = Inches(0.08)
            cell.margin_bottom = Inches(0.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = INDIGO if r == 0 else (WHITE if r % 2 else LIGHT)

            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            run = p.add_run()
            run.text = data[r][c]
            run.font.name = FONT
            run.font.size = Pt(15 if r == 0 else 13)
            run.font.bold = (r == 0 or c == 0)
            run.font.color.rgb = WHITE if r == 0 else (INDIGO if c == 0 else INK)

    add_text(
        slide,
        "Permissions are enforced on the server, not just hidden in the interface.",
        0.7, 6.35, 11.93, 0.35, size=13, italic=True, color=MUTED,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, 4)
    add_notes(
        slide,
        "Students own their complaints, admins own triage and oversight, and "
        "departments own the resolution work. Visibility is limited by role.",
    )

    # ---------------------------------------------------------------- slide 5
    slide = new_slide(prs)
    add_title(slide, "How It Works", "The complaint lifecycle")

    steps = [
        ("Filed", "Student"),
        ("Assigned", "Admin"),
        ("In Progress", "Department"),
        ("Resolved", "Department"),
        ("Closed", "Student"),
    ]
    box_w, box_h, arrow_w = 2.0, 0.85, 0.45
    top = 1.85
    x = 0.7
    for i, (label, actor) in enumerate(steps):
        box = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(top),
            Inches(box_w), Inches(box_h)
        )
        box.adjustments[0] = 0.12
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = INDIGO
        box.line.width = Pt(1.5)
        box.shadow.inherit = False
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.05)
        tf.margin_right = Inches(0.05)
        tf.margin_top = 0
        tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = label
        run.font.name = FONT
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = INDIGO

        add_text(slide, actor, x, top + 0.95, box_w, 0.3, size=10,
                 color=MUTED, align=PP_ALIGN.CENTER)

        if i < len(steps) - 1:
            arrow = slide.shapes.add_shape(
                MSO_SHAPE.RIGHT_ARROW, Inches(x + box_w + 0.05),
                Inches(top + 0.3), Inches(arrow_w - 0.1), Inches(0.25)
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = INDIGO
            arrow.line.fill.background()
            arrow.shadow.inherit = False
        x += box_w + arrow_w

    add_card(slide, 0.7, 3.55, 5.85, 0.8, fill=LIGHT, line=LIGHT)
    add_text(slide, "Reject: Resolved \u2192 In Progress", 0.95, 3.73, 5.4, 0.3,
             size=13, bold=True, color=INK)
    add_text(slide, "Student gives a mandatory reason", 0.95, 4.0, 5.4, 0.25,
             size=11, color=MUTED)

    add_card(slide, 6.78, 3.55, 5.85, 0.8, fill=LIGHT, line=LIGHT)
    add_text(slide, "Reopen: Closed \u2192 Reopened", 7.03, 3.73, 5.4, 0.3,
             size=13, bold=True, color=INK)
    add_text(slide, "Student reopens; priority rises one level", 7.03, 4.0, 5.4, 0.25,
             size=11, color=MUTED)

    add_card(slide, 0.7, 4.75, 11.93, 1.55, fill=WHITE)
    add_text(slide, "Under the hood", 0.95, 4.95, 5.0, 0.3, size=14, bold=True, color=INK)
    add_text(
        slide,
        "Django handles authentication, forms, and validation. Business rules live in a "
        "services layer, and every status change is written to a history table.",
        0.95, 5.35, 11.4, 0.7, size=14, color=INK,
    )
    add_footer(slide, 5)
    add_notes(
        slide,
        "The lifecycle is strict: Filed to Assigned to In Progress to Resolved to "
        "Closed. Students can reject a resolution or reopen a closed complaint, and "
        "every change is logged.",
    )

    # ---------------------------------------------------------------- slide 6
    slide = new_slide(prs)
    add_title(slide, "Tech Stack", "A Django application with a lightweight frontend")

    stack = [
        ("Backend", "Python 3.14  ·  Django 6.1.2"),
        ("Database", "SQLite  ·  Django ORM"),
        ("Frontend", "Django Templates  ·  Tailwind CSS  ·  vanilla JavaScript"),
        ("Authentication", "Django auth with email login"),
        ("Admin", "Django admin panel"),
        ("Quality", "Services layer  ·  database transactions  ·  audit history"),
        ("Testing", "55 automated tests with Django TestCase"),
    ]

    rows = len(stack) + 1
    table_shape = slide.shapes.add_table(
        rows, 2, Inches(0.7), Inches(1.65), Inches(11.93), Inches(4.85)
    )
    table = table_shape.table
    table.columns[0].width = Inches(2.7)
    table.columns[1].width = Inches(9.23)
    table.rows[0].height = Inches(0.55)
    for r in range(1, rows):
        table.rows[r].height = Inches(0.6)

    headers = ("Layer", "Technology")
    for c, header in enumerate(headers):
        cell = table.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = INDIGO
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = Inches(0.18)
        p = cell.text_frame.paragraphs[0]
        run = p.add_run()
        run.text = header
        run.font.name = FONT
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = WHITE

    for r, (layer, tech) in enumerate(stack, start=1):
        for c, value in enumerate((layer, tech)):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r % 2 else LIGHT
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.18)
            cell.margin_right = Inches(0.18)
            p = cell.text_frame.paragraphs[0]
            run = p.add_run()
            run.text = value
            run.font.name = FONT
            run.font.size = Pt(13)
            run.font.bold = (c == 0)
            run.font.color.rgb = INDIGO if c == 0 else INK

    add_footer(slide, 6)
    add_notes(
        slide,
        "The stack is deliberately simple: Django for the backend and templates, "
        "SQLite for storage, Tailwind for styling, and plain JavaScript for small "
        "interactions. There is no separate frontend framework.",
    )

    # ---------------------------------------------------------------- slide 7
    slide = new_slide(prs)
    add_title(slide, "Demo and Screens", "A quick look at the main screens")

    screens = ["Student Dashboard", "Complaint Detail", "Admin Dashboard"]
    box_w, gap = 3.75, 0.34
    x = 0.7
    for label in screens:
        box = add_card(slide, x, 1.65, box_w, 2.35, fill=LIGHT, line=INDIGO)
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.15)
        tf.margin_right = Inches(0.15)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = label + "\n"
        run.font.name = FONT
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = INDIGO
        run2 = p.add_run()
        run2.text = "Screenshot"
        run2.font.name = FONT
        run2.font.size = Pt(11)
        run2.font.color.rgb = MUTED
        x += box_w + gap

    add_card(slide, 0.7, 4.3, 11.93, 1.85, fill=WHITE)
    add_text(slide, "Demo accounts", 0.95, 4.5, 5.0, 0.3, size=15, bold=True, color=INK)
    accounts = [
        "Student       rajat@student.srmist.edu  /  student123",
        "Admin          admin@srmist.edu  /  admin123",
        "Department   mess@srmist.edu  /  dept123",
    ]
    y = 4.9
    for line in accounts:
        add_text(slide, line, 0.95, y, 8.5, 0.28, size=13, color=INK)
        y += 0.32
    add_text(slide, "Run:  python manage.py runserver", 9.1, 4.95, 3.3, 0.3,
             size=12, italic=True, color=MUTED, align=PP_ALIGN.RIGHT)

    add_footer(slide, 7)
    add_notes(
        slide,
        "Walk through one complaint as a student, assign it as an admin, resolve it "
        "as the department, then accept it back as the student. Demo accounts are "
        "listed on the slide.",
    )

    # ---------------------------------------------------------------- slide 8
    slide = new_slide(prs)
    add_title(slide, "Conclusion and Future Scope", "A small system with a complete workflow")

    left = add_card(slide, 0.7, 1.75, 5.85, 4.0, fill=WHITE)
    add_text(slide, "What It Delivers", 1.0, 1.98, 5.2, 0.4, size=18, bold=True, color=INK)
    add_bullets(
        slide,
        [
            "One tracked workflow for every complaint",
            "Clear ownership and status at all times",
            "Full audit trail of status changes",
            "A student verification loop",
            "Role-based access and visibility",
        ],
        1.0, 2.5, 5.25, 3.0, size=15, space_after=13,
    )

    right = add_card(slide, 6.78, 1.75, 5.85, 4.0, fill=LIGHT, line=LIGHT)
    add_text(slide, "Future Scope", 7.08, 1.98, 5.2, 0.4, size=18, bold=True, color=INDIGO)
    add_bullets(
        slide,
        [
            "Email and SMS notifications",
            "File attachments on complaints",
            "PostgreSQL and production deployment",
            "SLA timers and automatic escalation",
            "REST API and a mobile app",
        ],
        7.08, 2.5, 5.25, 3.0, size=15, space_after=13,
    )

    add_text(slide, "Thank you  ·  Questions?", 0.7, 6.15, 11.93, 0.5,
             size=24, bold=True, color=INDIGO, align=PP_ALIGN.CENTER)
    add_footer(slide, 8)
    add_notes(
        slide,
        "The project delivers a complete, accountable complaint workflow. Next steps "
        "are notifications, attachments, production deployment, and a mobile client.",
    )

    prs.save("CampusComplaint.pptx")
    print("Saved CampusComplaint.pptx")


if __name__ == "__main__":
    build()
